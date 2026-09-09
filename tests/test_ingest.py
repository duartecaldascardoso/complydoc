"""Loading each supported format, and failing safely on the ones we cannot."""

from __future__ import annotations

import pytest

from complydoc.discovery import discover
from complydoc.ingest.base import DocumentFormat, IngestOptions, LoaderError
from complydoc.ingest.registry import load_document, loader_for, supported_extensions
from tests.helpers import FIXTURES


@pytest.mark.parametrize(
    ("name", "expected"),
    [
        ("native_text.pdf", DocumentFormat.PDF),
        ("scan_page.png", DocumentFormat.IMAGE),
        ("sample.docx", DocumentFormat.DOCX),
        ("sample.xlsx", DocumentFormat.XLSX),
    ],
)
def test_each_format_loads(loader, name, expected):
    document = loader(name)
    assert document.format is expected
    assert document.pages
    assert len(document.sha256) == 64


def test_supported_extensions_cover_the_brief():
    extensions = set(supported_extensions())
    assert {".pdf", ".png", ".jpg", ".docx", ".xlsx"} <= extensions


def test_a_broken_file_raises_rather_than_crashing_the_run():
    with pytest.raises(LoaderError):
        load_document(FIXTURES / "broken.pdf", IngestOptions())


def test_discovery_reports_unsupported_files_instead_of_dying():
    files, skipped = discover(FIXTURES)
    assert files
    reasons = {s.path.name: s.reason for s in skipped}
    assert reasons.get("notes.txt") == "unsupported file type"
    assert loader_for(FIXTURES / "notes.txt") is None


def test_discovery_reports_a_missing_path():
    files, skipped = discover(FIXTURES / "nope")
    assert not files
    assert skipped and skipped[0].reason == "path does not exist"


def test_encrypted_pdf_is_flagged_and_not_fatal(loader):
    document = loader("encrypted.pdf")
    assert document.encrypted is True
    assert document.pages == []
    assert any("password protected" in w for w in document.load_warnings)


def test_scanned_page_has_no_text_layer(loader):
    document = loader("scanned_page.pdf")
    assert not document.has_text_layer
    page = document.pages[0]
    assert page.text_source == "none"
    assert page.image_blocks, "a scan should carry an embedded image"


def test_page_count_is_marked_unknown_for_flowing_formats(loader):
    assert loader("sample.docx").page_count_known is False
    assert loader("sample.xlsx").page_count_known is False
    assert loader("native_text.pdf").page_count_known is True


def test_a_page_with_no_text_layer_is_named_rather_than_ignored(loader):
    document = loader("scanned_page.pdf")
    joined = " ".join(document.load_warnings)
    assert "not read" in joined and "OCR" in joined


def test_worksheet_becomes_a_page(loader):
    document = loader("sample.xlsx")
    assert document.page_count == 2
    assert "Northwind" in document.full_text


def test_docx_merges_are_counted_from_the_markup(loader):
    """The count used to depend on which memory addresses got reused.

    Cells were identified by `id(cell._tc)`, and python-docx builds a new proxy
    on each access, so two different cells could share an address once the first
    had been collected. The fixture's one three-column header cell absorbs two
    grid positions and nothing else does.
    """
    tables = [t for page in loader("sample.docx").pages for t in page.tables]
    assert len(tables) == 1
    assert (tables[0].rows, tables[0].cols) == (4, 3)
    assert tables[0].merged_cells == 2
    assert tables[0].header_depth == 2


def test_docx_merges_count_the_same_in_a_fresh_process():
    """Reading the same file twice must not give two different answers."""
    import subprocess
    import sys

    script = (
        "from pathlib import Path;"
        "from complydoc.ingest.base import IngestOptions;"
        "from complydoc.ingest.registry import load_document;"
        "d = load_document(Path('tests/fixtures/sample.docx'), IngestOptions());"
        "print([t.merged_cells for p in d.pages for t in p.tables])"
    )
    runs = {
        subprocess.run(
            [sys.executable, "-c", script], capture_output=True, text=True, check=True
        ).stdout.strip()
        for _ in range(3)
    }
    assert runs == {"[2]"}, runs


def test_a_vertical_merge_is_counted_too(tmp_path):
    import docx

    from complydoc.ingest.docx import _table_info

    document = docx.Document()
    table = document.add_table(rows=4, cols=3)
    table.cell(0, 0).merge(table.cell(0, 2))
    table.cell(1, 0).merge(table.cell(2, 0))
    path = tmp_path / "merges.docx"
    document.save(path)

    info = _table_info(docx.Document(path).tables[0])
    # Twelve grid positions; the header cell absorbs two and the vertical one.
    assert info.merged_cells == 3
    assert (info.rows, info.cols) == (4, 3)
