"""Spreadsheet loader. Each worksheet is treated as one logical page."""

from __future__ import annotations

from pathlib import Path

import openpyxl
from openpyxl.utils.exceptions import InvalidFileException

from complydoc.ingest.base import (
    Document,
    DocumentFormat,
    IngestOptions,
    LoaderError,
    Page,
    TableInfo,
    sha256_of,
)
from complydoc.ingest.registry import register


class XlsxLoader:
    extensions = (".xlsx", ".xlsm")
    format = DocumentFormat.XLSX

    def load(self, path: Path, options: IngestOptions) -> Document:
        document = Document(
            path=path,
            sha256=sha256_of(path),
            format=self.format,
            page_count_known=False,
        )

        try:
            workbook = openpyxl.load_workbook(
                str(path), read_only=False, data_only=True, keep_links=False
            )
        except (InvalidFileException, OSError, KeyError, ValueError) as exc:
            raise LoaderError(f"openpyxl could not open the file: {exc}") from exc

        try:
            for index, sheet in enumerate(workbook.worksheets):
                page = Page(number=index + 1, width_pt=0.0, height_pt=0.0)
                rows = int(sheet.max_row or 0)
                cols = int(sheet.max_column or 0)

                merged_ranges = list(getattr(sheet, "merged_cells", []).ranges) or []
                merged_cells = sum(
                    (r.max_row - r.min_row + 1) * (r.max_col - r.min_col + 1) - 1
                    for r in merged_ranges
                )

                header_depth = 0
                for row_index in range(1, min(rows, 10) + 1):
                    if any(r.min_row <= row_index <= r.max_row for r in merged_ranges):
                        header_depth += 1
                    else:
                        break

                lines: list[str] = []
                for row in sheet.iter_rows(values_only=True):
                    values = [str(v) for v in row if v is not None and str(v).strip()]
                    if values:
                        lines.append("\t".join(values))

                page.text = "\n".join(lines)
                page.text_source = "native" if page.text.strip() else "none"
                if rows and cols:
                    page.tables.append(
                        TableInfo(
                            rows=rows,
                            cols=cols,
                            header_depth=max(1, header_depth),
                            merged_cells=merged_cells,
                        )
                    )
                page.notes.append(f"Worksheet {sheet.title!r} treated as one logical page.")
                document.pages.append(page)
        finally:
            workbook.close()

        document.load_warnings.append(
            "Printed page count, page dimensions and scan resolution are not determinable "
            "for a spreadsheet without rendering it, so they were not measured."
        )
        return document


register(XlsxLoader())
